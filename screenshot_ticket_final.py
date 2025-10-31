#!/usr/bin/env python3
"""
票务系统完整截图工具 - 最终版
遍历所有菜单和页面功能
"""

import os
import time
import json
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

class TicketSystemFinalScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_ticket_final"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_urls = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(3)
        
    def login(self):
        """登录"""
        print(f"\n正在登录票务系统...")
        self.driver.get(self.base_url)
        time.sleep(5)
        
        try:
            # 登录页截图
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面")
            
            username_input = self.driver.find_element(By.NAME, "LoginInput.UserNameOrEmailAddress")
            password_input = self.driver.find_element(By.NAME, "LoginInput.Password")
            
            username_input.clear()
            username_input.send_keys(self.username)
            password_input.clear()
            password_input.send_keys(self.password)
            
            # 使用Enter键提交
            password_input.send_keys(Keys.RETURN)
            print("✓ 已提交登录")
            
            # 等待登录完成
            for i in range(15):
                time.sleep(1)
                current_url = self.driver.current_url
                if 'login' not in current_url.lower():
                    print(f"✓ 登录成功！URL: {current_url}")
                    break
            
            time.sleep(3)
            
            # 主页截图
            self.save_screenshot("系统主页", "Dashboard", "登录后主页")
            
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            current_url = self.driver.current_url
            
            # 简单去重
            url_key = f"{current_url}_{function}"
            if url_key in self.visited_urls:
                return False
            
            self.visited_urls.add(url_key)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.8)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def get_main_menu_items(self):
        """获取主菜单项"""
        menu_items = []
        try:
            time.sleep(2)
            
            # 查找所有菜单链接
            links = self.driver.find_elements(By.XPATH, "//nav//a | //aside//a | //*[contains(@class, 'menu')]//a")
            
            for link in links:
                try:
                    if link.is_displayed():
                        text = link.text.strip()
                        href = link.get_attribute('href')
                        if text and len(text) < 100 and text not in ['EN', 'CN', 'ZH']:
                            menu_items.append({
                                'text': text,
                                'href': href,
                                'element': link
                            })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for item in menu_items:
                if item['text'] not in seen:
                    seen.add(item['text'])
                    unique.append(item)
            
            return unique
        except:
            return []
    
    def click_element(self, element):
        """点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
                return True
            except:
                self.driver.execute_script("arguments[0].click();", element)
                return True
        except:
            return False
    
    def find_page_buttons(self):
        """查找页面按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH, 
                "//button | //a[contains(@class, 'btn')] | //input[@type='submit']")
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text and len(text) < 50:
                            buttons.append({
                                'element': btn,
                                'text': text
                            })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen and btn['text'] not in ['EN', 'CN', 'ZH', 'English']:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique
        except:
            return []
    
    def explore_page(self, module_name, page_name):
        """探索页面功能"""
        print(f"    → 探索页面功能...")
        
        buttons = self.find_page_buttons()
        
        # 筛选重要按钮
        important_keywords = ['新增', '新建', 'Add', 'Create', 'New', '添加', '编辑', 'Edit', 
                             '查看', 'View', 'Details', '详情', '删除', 'Delete', 
                             '设置', 'Setting', 'Config', '导出', 'Export', '导入', 'Import',
                             '审核', 'Approve', '发布', 'Publish']
        
        important_btns = []
        for btn in buttons:
            text_lower = btn['text'].lower()
            if any(keyword.lower() in text_lower for keyword in important_keywords):
                important_btns.append(btn)
        
        print(f"      发现重要按钮: {len(important_btns)} 个")
        
        for i, btn_info in enumerate(important_btns[:6], 1):  # 限制6个
            try:
                btn_text = btn_info['text']
                print(f"        [{i}] 点击: {btn_text}")
                
                time.sleep(1)
                fresh_buttons = self.find_page_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element(target):
                    time.sleep(2)
                    
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{btn_text}功能")
                    
                    # 尝试返回
                    try:
                        self.driver.back()
                        time.sleep(2)
                    except:
                        pass
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def explore_all(self):
        """遍历所有菜单"""
        print("\n" + "=" * 80)
        print("开始全面遍历票务系统...")
        print("=" * 80)
        
        menus = self.get_main_menu_items()
        print(f"\n✓ 发现 {len(menus)} 个菜单项\n")
        
        for i, menu in enumerate(menus, 1):
            try:
                print(f"\n{'='*80}")
                print(f"[{i}/{len(menus)}] 菜单: {menu['text']}")
                print(f"{'='*80}")
                
                # 返回主页
                self.driver.get(self.base_url)
                time.sleep(2)
                
                # 重新查找菜单
                fresh_menus = self.get_main_menu_items()
                target = None
                for fm in fresh_menus:
                    if fm['text'] == menu['text']:
                        target = fm['element']
                        break
                
                if target and self.click_element(target):
                    time.sleep(3)
                    
                    # 截图
                    self.save_screenshot(menu['text'], f"{menu['text']}页面", f"{menu['text']}功能模块")
                    
                    # 探索页面功能
                    self.explore_page(menu['text'], menu['text'])
                
                print()
                
            except Exception as e:
                print(f"✗ 处理菜单出错: {str(e)}\n")
        
        print("=" * 80)
        print(f"✓ 遍历完成！共截图 {len(self.screenshots)} 个功能点")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word(screenshots, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading('票务系统 - 完整功能截图文档', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('票务系统（亿思维智慧门店）\n')
    info.add_run('系统地址：').bold = True
    info.add_run('https://xstest.axioxio.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y-%m-%d %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张')
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    header = table.rows[0].cells
    header[0].text = '序号'
    header[1].text = '功能点'
    header[2].text = '截图'
    
    for cell in header:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    widths = (Inches(0.5), Inches(2.5), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    for ss in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        cells[0].text = str(ss['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"【{ss['module']}】\n")
        run.font.bold = True
        
        run = desc_para.add_run(f"{ss['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{ss['description']}")
        run.font.size = Pt(8)
        
        img_path = os.path.join('/workspace/screenshots_ticket_final', ss['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片失败]')
    
    doc.save(output_path)
    print(f"✓ Word文档: {output_path} ({os.path.getsize(output_path) / 1024 / 1024:.2f} MB)")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    
    print("=" * 80)
    print("票务系统完整截图工具")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = TicketSystemFinalScreenshotter(base_url, username, password)
        
        if screenshotter.login():
            screenshotter.explore_all()
            
            # 保存JSON
            json_path = os.path.join('screenshots_ticket_final', 'index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'total': len(screenshotter.screenshots),
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            # 生成Word
            doc_path = '/workspace/票务系统完整功能截图文档（最终版）.docx'
            create_word(screenshotter.screenshots, doc_path)
            
            # 打包
            print("\n正在打包...")
            os.system("cd /workspace && zip -q -r 票务系统完整功能截图包（最终版）.zip screenshots_ticket_final/ 票务系统完整功能截图文档（最终版）.docx")
            
            print("\n" + "=" * 80)
            print("✅ 完成！")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统完整功能截图包（最终版）.zip")
            print("=" * 80)
        else:
            print("\n❌ 登录失败")
    
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
