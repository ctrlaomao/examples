#!/usr/bin/env python3
"""
票务系统完整截图工具
遍历所有级别菜单和页面功能点
"""

import os
import time
import json
import hashlib
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import TimeoutException, NoSuchElementException
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

class TicketSystemScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_ticket_system"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_hashes = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--disable-gpu')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(3)
        
    def get_page_hash(self):
        """获取页面内容哈希"""
        try:
            url = self.driver.current_url
            body = self.driver.find_element(By.TAG_NAME, "body").text[:1000]
            return hashlib.md5(f"{url}|{body}".encode()).hexdigest()
        except:
            return hashlib.md5(str(time.time()).encode()).hexdigest()
    
    def login(self):
        """登录系统"""
        print(f"\n正在访问票务系统: {self.base_url}")
        self.driver.get(self.base_url)
        time.sleep(4)
        
        try:
            # 保存登录页截图
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面")
            
            # 查找用户名和密码输入框
            username_selectors = [
                (By.NAME, "username"),
                (By.NAME, "phone"),
                (By.NAME, "mobile"),
                (By.ID, "username"),
                (By.XPATH, "//input[@type='text']"),
                (By.XPATH, "//input[@placeholder='手机号']"),
                (By.XPATH, "//input[@placeholder='用户名']"),
            ]
            
            password_selectors = [
                (By.NAME, "password"),
                (By.NAME, "pwd"),
                (By.ID, "password"),
                (By.XPATH, "//input[@type='password']"),
            ]
            
            username_input = None
            for selector_type, selector_value in username_selectors:
                try:
                    username_input = self.driver.find_element(selector_type, selector_value)
                    print(f"✓ 找到用户名输入框: {selector_value}")
                    break
                except:
                    pass
            
            password_input = None
            for selector_type, selector_value in password_selectors:
                try:
                    password_input = self.driver.find_element(selector_type, selector_value)
                    print(f"✓ 找到密码输入框: {selector_value}")
                    break
                except:
                    pass
            
            if not username_input or not password_input:
                print("✗ 无法找到登录表单")
                return False
            
            # 输入用户名和密码
            username_input.clear()
            username_input.send_keys(self.username)
            password_input.clear()
            password_input.send_keys(self.password)
            print(f"✓ 已输入用户名和密码")
            
            # 查找登录按钮
            login_button_selectors = [
                (By.XPATH, "//button[contains(text(), '登录')]"),
                (By.XPATH, "//button[contains(text(), '登錄')]"),
                (By.XPATH, "//button[@type='submit']"),
                (By.XPATH, "//button[contains(@class, 'login')]"),
                (By.XPATH, "//a[contains(text(), '登录')]"),
            ]
            
            login_button = None
            for selector_type, selector_value in login_button_selectors:
                try:
                    login_button = self.driver.find_element(selector_type, selector_value)
                    print(f"✓ 找到登录按钮")
                    break
                except:
                    pass
            
            if login_button:
                login_button.click()
            else:
                password_input.submit()
            
            print("✓ 已提交登录")
            
            # 等待登录完成
            for i in range(15):
                time.sleep(1)
                current_url = self.driver.current_url
                if current_url != self.base_url and 'login' not in current_url.lower():
                    print(f"✓ 登录成功！当前URL: {current_url}")
                    break
                print(f"  等待登录... ({i+1}/15)")
            
            time.sleep(3)
            
            # 保存登录后主页
            self.save_screenshot("系统主页", "主页Dashboard", "登录后主页")
            
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            import traceback
            traceback.print_exc()
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            page_hash = self.get_page_hash()
            if page_hash in self.visited_hashes:
                return False
            
            self.visited_hashes.add(page_hash)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.5)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def get_all_menus(self):
        """获取所有菜单项"""
        menus = []
        try:
            time.sleep(2)
            
            # 多种菜单选择器
            menu_selectors = [
                "//li[contains(@class, 'el-submenu') or contains(@class, 'el-menu-item')]",
                "//div[contains(@class, 'menu')]//li",
                "//nav//li",
                "//aside//li",
                "//*[contains(@class, 'sidebar')]//li",
                "//*[@role='menuitem']",
            ]
            
            all_elements = []
            for selector in menu_selectors:
                try:
                    elements = self.driver.find_elements(By.XPATH, selector)
                    all_elements.extend(elements)
                except:
                    pass
            
            seen = set()
            for elem in all_elements:
                try:
                    if elem.is_displayed():
                        text = elem.text.strip().split('\n')[0]
                        if text and len(text) < 50 and text not in seen:
                            seen.add(text)
                            class_attr = elem.get_attribute('class') or ''
                            is_submenu = 'submenu' in class_attr.lower()
                            menus.append({
                                'text': text,
                                'element': elem,
                                'is_submenu': is_submenu
                            })
                except:
                    pass
            
            return menus
        except Exception as e:
            print(f"  获取菜单失败: {str(e)}")
            return []
    
    def get_submenus(self, parent_element):
        """获取子菜单"""
        submenus = []
        try:
            time.sleep(1)
            
            # 查找子菜单的多种方式
            sub_selectors = [
                ".//ul//li",
                ".//div[contains(@class, 'submenu')]//li",
                ".//li",
            ]
            
            all_subs = []
            for selector in sub_selectors:
                try:
                    subs = parent_element.find_elements(By.XPATH, selector)
                    all_subs.extend(subs)
                except:
                    pass
            
            for elem in all_subs:
                try:
                    if elem.is_displayed():
                        text = elem.text.strip()
                        if text and len(text) < 50:
                            submenus.append({
                                'text': text,
                                'element': elem
                            })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for item in submenus:
                if item['text'] not in seen:
                    seen.add(item['text'])
                    unique.append(item)
            
            return unique
        except:
            return []
    
    def click_element(self, element):
        """点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
                return True
            except:
                self.driver.execute_script("arguments[0].click();", element)
                return True
        except:
            return False
    
    def find_all_buttons(self):
        """查找页面所有按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH,
                "//button | //a[contains(@class, 'button') or contains(@class, 'btn')]")
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text:
                            buttons.append({
                                'element': btn,
                                'text': text
                            })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique
        except:
            return []
    
    def close_dialog(self):
        """关闭弹窗"""
        try:
            close_selectors = [
                "//*[contains(@class, 'close')]",
                "//*[contains(@class, 'el-dialog__close')]",
                "//*[contains(@class, 'modal-close')]",
                "//button[contains(text(), '取消')]",
                "//button[contains(text(), '关闭')]",
            ]
            
            for selector in close_selectors:
                try:
                    close_btns = self.driver.find_elements(By.XPATH, selector)
                    for btn in close_btns:
                        if btn.is_displayed():
                            self.click_element(btn)
                            time.sleep(0.5)
                            return True
                except:
                    pass
            return False
        except:
            return False
    
    def explore_page_functions(self, module_name, page_name):
        """探索页面功能"""
        print(f"    → 探索页面功能: {page_name}")
        
        buttons = self.find_all_buttons()
        
        # 重要操作关键词
        important_keywords = ['新增', '新建', '添加', '创建', '编辑', '修改', '查看', '详情',
                             '删除', '设置', '配置', '管理', '导出', '导入', '上传', '下载',
                             '审核', '发布', '上架', '下架', '启用', '禁用', '查询', '搜索']
        
        important_buttons = []
        for btn in buttons:
            if any(keyword in btn['text'] for keyword in important_keywords):
                important_buttons.append(btn)
        
        print(f"      发现重要操作按钮: {len(important_buttons)} 个")
        
        # 处理按钮
        for i, btn_info in enumerate(important_buttons[:8]):  # 限制8个
            try:
                btn_text = btn_info['text']
                print(f"        [{i+1}] 点击: {btn_text}")
                
                time.sleep(0.8)
                fresh_buttons = self.find_all_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element(target):
                    time.sleep(2)
                    
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{btn_text}功能页面")
                    
                    if self.close_dialog():
                        time.sleep(0.5)
                    else:
                        try:
                            self.driver.back()
                            time.sleep(1.5)
                        except:
                            pass
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def explore_all_system(self):
        """遍历整个系统"""
        print("\n" + "=" * 80)
        print("开始全面遍历票务系统...")
        print("=" * 80)
        
        menus = self.get_all_menus()
        print(f"\n✓ 发现 {len(menus)} 个菜单项\n")
        
        for i, menu_item in enumerate(menus, 1):
            try:
                print(f"\n{'='*80}")
                print(f"[{i}/{len(menus)}] 菜单: {menu_item['text']}")
                print(f"{'='*80}")
                
                time.sleep(1)
                fresh_menus = self.get_all_menus()
                target = None
                for fm in fresh_menus:
                    if fm['text'] == menu_item['text']:
                        target = fm
                        break
                
                if not target:
                    print(f"  ⚠ 无法找到菜单项")
                    continue
                
                if self.click_element(target['element']):
                    time.sleep(2)
                    
                    if target['is_submenu']:
                        # 有子菜单
                        submenus = self.get_submenus(target['element'])
                        print(f"  → 包含 {len(submenus)} 个子菜单\n")
                        
                        for j, submenu in enumerate(submenus, 1):
                            try:
                                print(f"  [{j}/{len(submenus)}] 子菜单: {submenu['text']}")
                                
                                time.sleep(1)
                                fresh_menus2 = self.get_all_menus()
                                fresh_parent = None
                                for fm2 in fresh_menus2:
                                    if fm2['text'] == menu_item['text']:
                                        fresh_parent = fm2['element']
                                        break
                                
                                if fresh_parent:
                                    fresh_subs = self.get_submenus(fresh_parent)
                                    for fs in fresh_subs:
                                        if fs['text'] == submenu['text']:
                                            if self.click_element(fs['element']):
                                                time.sleep(2)
                                                
                                                page_name = submenu['text']
                                                self.save_screenshot(menu_item['text'], page_name, f"{page_name}页面")
                                                
                                                self.explore_page_functions(menu_item['text'], page_name)
                                            break
                                
                                print()
                            except Exception as e:
                                print(f"  ✗ 处理子菜单出错: {str(e)}\n")
                    else:
                        # 无子菜单
                        page_name = menu_item['text']
                        self.save_screenshot(menu_item['text'], page_name, f"{page_name}页面")
                        
                        self.explore_page_functions(menu_item['text'], page_name)
                
                print()
                
            except Exception as e:
                print(f"✗ 处理菜单出错: {str(e)}\n")
        
        print("=" * 80)
        print(f"✓ 遍历完成！共截图 {len(self.screenshots)} 个功能点")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, output_path, system_name="票务系统"):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    print("=" * 80)
    
    doc = Document()
    
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading(f'{system_name} - 完整功能截图文档', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run(f'{system_name}\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y-%m-%d %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张')
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    header = table.rows[0].cells
    header[0].text = '序号'
    header[1].text = '功能点'
    header[2].text = '截图'
    
    for cell in header:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(11)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    widths = (Inches(0.5), Inches(2.5), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    for ss in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        cells[0].text = str(ss['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"【{ss['module']}】\n")
        run.font.bold = True
        run.font.size = Pt(10)
        
        run = desc_para.add_run(f"{ss['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{ss['description']}")
        run.font.size = Pt(8)
        
        img_path = os.path.join('/workspace/screenshots_ticket_system', ss['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片加载失败]')
        else:
            cells[2].text = '[图片不存在]'
    
    doc.save(output_path)
    
    print("=" * 80)
    print(f"✓ Word文档生成完成")
    print(f"📄 文件: {output_path}")
    print(f"📊 大小: {os.path.getsize(output_path) / 1024 / 1024:.2f} MB")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    output_dir = "screenshots_ticket_system"
    
    print("=" * 80)
    print("票务系统完整截图工具")
    print("=" * 80)
    print(f"目标系统: {base_url}")
    print(f"用户名: {username}")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = TicketSystemScreenshotter(base_url, username, password, output_dir)
        
        if screenshotter.login():
            screenshotter.explore_all_system()
            
            # 保存JSON索引
            json_path = os.path.join(output_dir, 'screenshots_index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'system': '票务系统',
                    'url': base_url,
                    'total': len(screenshotter.screenshots),
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            # 生成Word文档
            doc_path = '/workspace/票务系统完整功能截图文档.docx'
            create_word_document(screenshotter.screenshots, doc_path, "票务系统")
            
            # 打包
            print("\n正在打包所有文件...")
            os.system(f"cd /workspace && zip -q -r 票务系统完整功能截图包.zip {output_dir}/ 票务系统完整功能截图文档.docx")
            
            print("\n" + "=" * 80)
            print("✅ 所有任务完成！")
            print(f"📁 截图目录: /workspace/{output_dir}")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统完整功能截图包.zip")
            print("=" * 80)
        else:
            print("\n❌ 登录失败，无法继续")
    
    except KeyboardInterrupt:
        print("\n\n⚠️  用户中断")
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
