#!/usr/bin/env python3
"""
票务系统完整截图工具 - 中文版
遍历所有菜单、子菜单和功能点
"""

import os
import time
import json
import hashlib
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


class TicketSystemChineseScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_ticket_chinese"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_states = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        chrome_options.add_argument('--lang=zh-CN')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(2)
    
    def get_page_state_hash(self):
        """获取页面状态哈希（用于去重）"""
        try:
            url = self.driver.current_url
            title = self.driver.title
            # 获取主内容区域的文本
            try:
                main_text = self.driver.find_element(By.TAG_NAME, "main").text[:300]
            except:
                main_text = self.driver.find_element(By.TAG_NAME, "body").text[:300]
            
            state = f"{url}_{title}_{main_text}"
            return hashlib.md5(state.encode()).hexdigest()
        except:
            return None
    
    def switch_to_chinese(self):
        """切换到简体中文"""
        try:
            print("  → 切换到简体中文...")
            
            # 点击语言下拉按钮
            lang_button = self.driver.find_element(By.XPATH, "//button[contains(@class, 'dropdown-toggle')]")
            lang_button.click()
            time.sleep(1)
            
            # 点击简体中文选项
            chinese_link = self.driver.find_element(By.XPATH, "//a[contains(text(), '简体中文')]")
            chinese_link.click()
            time.sleep(3)
            
            print("  ✓ 已切换到简体中文")
            return True
        except Exception as e:
            print(f"  ✗ 切换语言失败: {str(e)}")
            return False
    
    def login(self):
        """登录"""
        print(f"\n正在登录票务系统...")
        self.driver.get(self.base_url)
        time.sleep(5)
        
        try:
            # 切换语言
            self.switch_to_chinese()
            
            # 登录页截图
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面（中文）")
            
            username_input = self.driver.find_element(By.NAME, "LoginInput.UserNameOrEmailAddress")
            password_input = self.driver.find_element(By.NAME, "LoginInput.Password")
            
            username_input.clear()
            username_input.send_keys(self.username)
            password_input.clear()
            password_input.send_keys(self.password)
            password_input.send_keys(Keys.RETURN)
            
            print("✓ 已提交登录")
            
            # 等待登录完成
            for i in range(15):
                time.sleep(1)
                current_url = self.driver.current_url
                if 'login' not in current_url.lower():
                    print(f"✓ 登录成功！")
                    break
            
            time.sleep(3)
            self.save_screenshot("系统主页", "首页Dashboard", "登录后主页")
            
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            state_hash = self.get_page_state_hash()
            if state_hash and state_hash in self.visited_states:
                return False
            
            if state_hash:
                self.visited_states.add(state_hash)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.8)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def click_element_safe(self, element):
        """安全点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
                return True
            except:
                self.driver.execute_script("arguments[0].click();", element)
                return True
        except:
            return False
    
    def close_dialog(self):
        """尝试关闭弹窗"""
        try:
            close_selectors = [
                "//button[contains(@class, 'close')]",
                "//button[contains(text(), '取消')]",
                "//button[contains(text(), '关闭')]",
                "//a[contains(text(), '返回')]",
                "//button[contains(@class, 'ant-modal-close')]",
                "//button[contains(@class, 'el-dialog__close')]"
            ]
            
            for selector in close_selectors:
                try:
                    close_btns = self.driver.find_elements(By.XPATH, selector)
                    for btn in close_btns:
                        if btn.is_displayed():
                            btn.click()
                            time.sleep(1)
                            return True
                except:
                    pass
            
            # 尝试按ESC键
            from selenium.webdriver.common.action_chains import ActionChains
            ActionChains(self.driver).send_keys(Keys.ESCAPE).perform()
            time.sleep(1)
        except:
            pass
    
    def find_page_buttons(self):
        """查找页面重要按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH, 
                "//button | //a[contains(@class, 'btn')] | //input[@type='submit'] | //input[@type='button']")
            
            important_keywords = ['新增', '新建', '添加', '编辑', '修改', '查看', '详情', 
                                 '删除', '设置', '配置', '导出', '导入', '审核', '发布', 
                                 '提交', '保存', '搜索', '查询', '筛选', '打印']
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text and len(text) < 50:
                            text_lower = text.lower()
                            if any(keyword in text for keyword in important_keywords):
                                buttons.append({
                                    'element': btn,
                                    'text': text
                                })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique[:10]  # 限制最多10个
        except:
            return []
    
    def explore_page_functions(self, module_name, page_name):
        """探索页面功能点"""
        buttons = self.find_page_buttons()
        
        if buttons:
            print(f"      发现功能按钮: {len(buttons)} 个")
        
        for i, btn_info in enumerate(buttons, 1):
            try:
                btn_text = btn_info['text']
                print(f"        [{i}] 点击: {btn_text}")
                
                time.sleep(1)
                
                # 重新查找按钮（防止DOM变化）
                fresh_buttons = self.find_page_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element_safe(target):
                    time.sleep(2.5)
                    
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{page_name}的{btn_text}功能")
                    
                    # 关闭弹窗
                    self.close_dialog()
                    time.sleep(1)
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def get_all_submenu_items(self):
        """获取所有子菜单项（强制展开）"""
        try:
            # 使用JavaScript强制展开所有子菜单
            self.driver.execute_script("""
                var innerMenus = document.querySelectorAll('.lpx-inner-menu');
                innerMenus.forEach(function(menu) {
                    menu.classList.remove('collapsed');
                    menu.style.display = 'block';
                });
            """)
            time.sleep(1)
            
            # 查找所有子菜单链接
            sub_links = self.driver.find_elements(By.XPATH, 
                "//ul[contains(@class, 'lpx-inner-menu')]//a")
            
            submenus = []
            for link in sub_links:
                try:
                    text = link.text.strip()
                    href = link.get_attribute('href')
                    if text and len(text) < 100:
                        submenus.append({
                            'text': text,
                            'href': href
                        })
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for item in submenus:
                if item['text'] not in seen:
                    seen.add(item['text'])
                    unique.append(item)
            
            return unique
        except:
            return []
    
    def explore_all_menus(self):
        """遍历所有菜单"""
        print("\n" + "=" * 80)
        print("开始全面遍历票务系统（中文版）...")
        print("=" * 80)
        
        # 定义所有菜单结构（基于之前的检查结果）
        menu_structure = {
            "首页": [],
            "票务": [
                "收银", "发卡", "还卡", "订单", "票务管理",
                "电子券列表", "初始化", "纸制票模板",
                "人脸库", "闸机", "软件版本", "卡片管理", "通行记录", "通行规则",
                "产品管理", "产品分类"
            ],
            "预约": [
                "团队下单"
            ],
            "租赁": [],
            "分销商": [],
            "小程序设置": [],
            "销售网点": [],
            "报表": [],
            "系统管理": []
        }
        
        total_processed = 0
        
        for parent_menu, submenus in menu_structure.items():
            try:
                print(f"\n{'='*80}")
                print(f"一级菜单: {parent_menu}")
                print(f"{'='*80}")
                
                # 返回主页
                self.driver.get(self.base_url)
                time.sleep(3)
                
                # 展开所有子菜单
                self.get_all_submenu_items()
                
                # 点击一级菜单
                try:
                    menu_link = self.driver.find_element(By.XPATH, 
                        f"//aside//a[contains(text(), '{parent_menu}')] | //nav//a[contains(text(), '{parent_menu}')]")
                    
                    if self.click_element_safe(menu_link):
                        time.sleep(2)
                        
                        if not submenus:
                            # 没有子菜单，直接截图
                            print(f"  → 截图: {parent_menu}")
                            self.save_screenshot(parent_menu, f"{parent_menu}页面", f"{parent_menu}功能模块")
                            self.explore_page_functions(parent_menu, parent_menu)
                            total_processed += 1
                        else:
                            # 有子菜单，遍历子菜单
                            print(f"  → 发现 {len(submenus)} 个子菜单")
                            
                            for submenu in submenus:
                                try:
                                    print(f"\n  子菜单: {submenu}")
                                    
                                    # 返回主页重新导航
                                    self.driver.get(self.base_url)
                                    time.sleep(2)
                                    
                                    # 展开子菜单
                                    self.get_all_submenu_items()
                                    
                                    # 点击一级菜单
                                    menu_link = self.driver.find_element(By.XPATH, 
                                        f"//aside//a[contains(text(), '{parent_menu}')] | //nav//a[contains(text(), '{parent_menu}')]")
                                    self.click_element_safe(menu_link)
                                    time.sleep(2)
                                    
                                    # 展开子菜单
                                    self.get_all_submenu_items()
                                    
                                    # 点击子菜单
                                    submenu_link = self.driver.find_element(By.XPATH, 
                                        f"//ul[contains(@class, 'lpx-inner-menu')]//a[contains(text(), '{submenu}')]")
                                    
                                    if self.click_element_safe(submenu_link):
                                        time.sleep(3)
                                        
                                        # 截图子菜单页面
                                        self.save_screenshot(parent_menu, submenu, f"{parent_menu} - {submenu}")
                                        
                                        # 探索页面功能
                                        self.explore_page_functions(parent_menu, submenu)
                                        total_processed += 1
                                    
                                except Exception as e:
                                    print(f"  ✗ 处理子菜单 {submenu} 出错: {str(e)}")
                
                except Exception as e:
                    print(f"✗ 未找到菜单 {parent_menu}: {str(e)}")
            
            except Exception as e:
                print(f"✗ 处理菜单 {parent_menu} 出错: {str(e)}")
        
        print("\n" + "=" * 80)
        print(f"✓ 遍历完成！共截图 {len(self.screenshots)} 个功能点")
        print(f"✓ 处理了 {total_processed} 个页面")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading('票务系统 - 完整功能截图文档（中文版）', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('票务系统（亿思维智慧门店）\n')
    info.add_run('系统地址：').bold = True
    info.add_run('https://xstest.axioxio.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y年%m月%d日 %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张\n')
    info.add_run('语言版本：').bold = True
    info.add_run('简体中文')
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    header = table.rows[0].cells
    header[0].text = '序号'
    header[1].text = '功能模块 / 功能点'
    header[2].text = '截图'
    
    for cell in header:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(11)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    widths = (Inches(0.5), Inches(2.5), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    for ss in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        cells[0].text = str(ss['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"【{ss['module']}】\n")
        run.font.bold = True
        run.font.size = Pt(10)
        
        run = desc_para.add_run(f"{ss['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{ss['description']}")
        run.font.size = Pt(8)
        
        img_path = os.path.join('/workspace/screenshots_ticket_chinese', ss['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片加载失败]')
    
    doc.save(output_path)
    file_size = os.path.getsize(output_path) / 1024 / 1024
    print(f"✓ Word文档: {output_path} ({file_size:.2f} MB)")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    
    print("=" * 80)
    print("票务系统完整截图工具（中文版）")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = TicketSystemChineseScreenshotter(base_url, username, password)
        
        if screenshotter.login():
            screenshotter.explore_all_menus()
            
            # 保存JSON
            json_path = os.path.join('screenshots_ticket_chinese', 'index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'total': len(screenshotter.screenshots),
                    'language': '简体中文',
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            print(f"\n✓ JSON索引: {json_path}")
            
            # 生成Word
            doc_path = '/workspace/票务系统完整功能截图文档（中文版）.docx'
            create_word_document(screenshotter.screenshots, doc_path)
            
            # 打包
            print("\n正在打包...")
            os.system("cd /workspace && rm -f 票务系统完整功能截图包（中文版）.zip && " +
                     "zip -q -r 票务系统完整功能截图包（中文版）.zip screenshots_ticket_chinese/ 票务系统完整功能截图文档（中文版）.docx")
            
            zip_size = os.path.getsize('/workspace/票务系统完整功能截图包（中文版）.zip') / 1024 / 1024
            
            print("\n" + "=" * 80)
            print("✅ 完成！")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"🌐 语言版本: 简体中文")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统完整功能截图包（中文版）.zip ({zip_size:.2f} MB)")
            print("=" * 80)
        else:
            print("\n❌ 登录失败")
    
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
