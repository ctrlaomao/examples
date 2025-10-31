#!/usr/bin/env python3
"""
生成简单的Word文档 - 功能点对照列表
"""

import os
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import time

def create_simple_document():
    """创建简单的功能点对照文档"""
    
    # 读取报告数据
    with open('/workspace/screenshots_full/report.json', 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    screenshots = data['screenshots']
    
    # 创建Word文档
    doc = Document()
    
    # 设置默认中文字体
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    # 添加标题
    title = doc.add_heading('海星育后台管理系统 - 功能截图对照表', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 添加基本信息
    info = doc.add_paragraph()
    info.add_run('系统地址：').bold = True
    info.add_run('https://jhtest.bjstarfish.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y-%m-%d %H:%M:%S")}\n')
    info.add_run('总功能数：').bold = True
    info.add_run(f'{len(screenshots)} 个')
    
    doc.add_paragraph()  # 空行
    
    # 创建表格
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    # 设置表头
    header_cells = table.rows[0].cells
    header_cells[0].text = '功能点描述'
    header_cells[1].text = '功能截图'
    
    # 表头加粗居中
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 设置列宽
    table.autofit = False
    table.allow_autofit = False
    widths = (Inches(2.5), Inches(4.5))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    # 添加数据行
    print(f"\n正在生成Word文档...")
    print("=" * 60)
    
    for i, screenshot in enumerate(screenshots, 1):
        print(f"  [{i}/{len(screenshots)}] {screenshot['name']}")
        
        # 添加新行
        row = table.add_row()
        cells = row.cells
        
        # 设置列宽
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        # 第一列：功能点描述
        desc_para = cells[0].paragraphs[0]
        
        # 序号和功能名
        run = desc_para.add_run(f"{screenshot['index']}. {screenshot['name']}\n\n")
        run.font.size = Pt(10)
        run.font.bold = True
        
        # 页面标题
        run = desc_para.add_run(f"【页面】{screenshot['title']}\n")
        run.font.size = Pt(9)
        
        # URL
        run = desc_para.add_run(f"【链接】{screenshot['url']}")
        run.font.size = Pt(8)
        
        # 第二列：截图
        img_path = os.path.join('/workspace/screenshots_full', screenshot['filename'])
        
        if os.path.exists(img_path):
            img_para = cells[1].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            try:
                # 添加图片，设置合适的宽度
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(4.2))
            except Exception as e:
                img_para.add_run(f"[图片加载失败]")
                print(f"    ✗ 图片加载失败: {str(e)}")
        else:
            cells[1].text = '[图片文件不存在]'
            print(f"    ✗ 图片文件不存在")
    
    # 保存文档
    output_path = '/workspace/功能截图对照表.docx'
    doc.save(output_path)
    
    print("=" * 60)
    print(f"\n✅ Word文档生成成功！")
    print(f"📁 文件位置: {output_path}")
    print(f"📊 包含功能: {len(screenshots)} 个")
    print(f"📄 文件大小: {os.path.getsize(output_path) / 1024 / 1024:.2f} MB")
    
    return output_path


if __name__ == '__main__':
    create_simple_document()
