#!/usr/bin/env python3
"""
票务系统功能模块截图工具
按照功能清单组织截图
"""

import os
import time
import json
import hashlib
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


class FunctionModuleScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_by_modules"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_states = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        chrome_options.add_argument('--lang=zh-CN')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(3)
        
        # 功能模块定义
        self.function_modules = {
            "票务管理": {
                "description": "实现对雪票电子或纸质的全面管理",
                "menu_keywords": ["票务", "电子券", "纸制票", "票务管理", "初始化", "模板"],
                "subpages": []
            },
            "订单管理": {
                "description": "对各类雪票、教务订单、租赁订单进行管理",
                "menu_keywords": ["订单", "Orders"],
                "subpages": []
            },
            "工作台": {
                "description": "实现对现场售票及教练预约的下单管理",
                "menu_keywords": ["工作台", "收银", "发卡", "还卡", "团队下单"],
                "subpages": []
            },
            "通行管理": {
                "description": "通过闸机的管控实现对进出雪场的人员进行管理；通行方式可以是人脸或雪票二维码",
                "menu_keywords": ["闸机", "人脸库", "通行记录", "通行规则", "卡片管理"],
                "subpages": []
            },
            "控制面板": {
                "description": "包括对雪票产品设置，电子/纸质票设置、闸机配置、售票小程序、教练、分销商、节假日、租赁物等运营参数的管理",
                "menu_keywords": ["产品", "产品管理", "产品分类", "小程序", "分销商", "租赁"],
                "subpages": []
            },
            "报表管理": {
                "description": "提供各类运营报表，包括票务核销报表、教学核销报表、教学数据看板等",
                "menu_keywords": ["报表", "Reports"],
                "subpages": []
            },
            "系统管理": {
                "description": "系统基础功能的设置，包括文件，身份标识，日志及相关系统参数的设置",
                "menu_keywords": ["管理", "系统", "Administration"],
                "subpages": []
            }
        }
    
    def get_page_state_hash(self):
        """获取页面状态哈希"""
        try:
            url = self.driver.current_url
            title = self.driver.title
            try:
                main_text = self.driver.find_element(By.TAG_NAME, "main").text[:200]
            except:
                main_text = self.driver.find_element(By.TAG_NAME, "body").text[:200]
            state = f"{url}_{title}_{main_text}"
            return hashlib.md5(state.encode()).hexdigest()
        except:
            return None
    
    def switch_to_chinese(self):
        """切换到简体中文"""
        try:
            print("  → 切换到简体中文...")
            lang_button = self.driver.find_element(By.XPATH, "//button[contains(@class, 'dropdown-toggle')]")
            lang_button.click()
            time.sleep(1)
            chinese_link = self.driver.find_element(By.XPATH, "//a[contains(text(), '简体中文')]")
            chinese_link.click()
            time.sleep(3)
            print("  ✓ 已切换到简体中文")
            return True
        except:
            return False
    
    def login(self):
        """登录"""
        print(f"\n正在登录票务系统...")
        self.driver.get(self.base_url)
        time.sleep(5)
        
        try:
            self.switch_to_chinese()
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面（中文）")
            
            username_input = self.driver.find_element(By.NAME, "LoginInput.UserNameOrEmailAddress")
            password_input = self.driver.find_element(By.NAME, "LoginInput.Password")
            
            username_input.send_keys(self.username)
            password_input.send_keys(self.password)
            password_input.send_keys(Keys.RETURN)
            
            print("✓ 已提交登录")
            
            for i in range(15):
                time.sleep(1)
                if 'login' not in self.driver.current_url.lower():
                    print(f"✓ 登录成功！")
                    break
            
            time.sleep(3)
            self.save_screenshot("系统主页", "首页Dashboard", "登录后主页")
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            state_hash = self.get_page_state_hash()
            if state_hash and state_hash in self.visited_states:
                return False
            if state_hash:
                self.visited_states.add(state_hash)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.8)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def click_element_safe(self, element):
        """安全点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
            except:
                self.driver.execute_script("arguments[0].click();", element)
            return True
        except:
            return False
    
    def close_dialog(self):
        """关闭弹窗"""
        try:
            for selector in ["//button[contains(text(), '取消')]", "//button[contains(text(), '关闭')]",
                           "//button[contains(@class, 'close')]"]:
                try:
                    btns = self.driver.find_elements(By.XPATH, selector)
                    for btn in btns:
                        if btn.is_displayed():
                            btn.click()
                            time.sleep(1)
                            return
                except:
                    pass
            from selenium.webdriver.common.action_chains import ActionChains
            ActionChains(self.driver).send_keys(Keys.ESCAPE).perform()
            time.sleep(1)
        except:
            pass
    
    def find_page_buttons(self):
        """查找页面按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH, "//button | //a[contains(@class, 'btn')]")
            
            keywords = ['新增', '新建', '添加', '编辑', '修改', '查看', '详情', 
                       '删除', '设置', '配置', '导出', '导入', '审核', '发布', 
                       '提交', '保存', '搜索', '查询', '筛选']
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text and len(text) < 50 and any(k in text for k in keywords):
                            buttons.append({'element': btn, 'text': text})
                except:
                    pass
            
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique[:10]
        except:
            return []
    
    def explore_page(self, module_name, page_name):
        """探索页面功能"""
        buttons = self.find_page_buttons()
        if buttons:
            print(f"      发现功能按钮: {len(buttons)} 个")
        
        for i, btn_info in enumerate(buttons, 1):
            try:
                btn_text = btn_info['text']
                print(f"        [{i}] 点击: {btn_text}")
                time.sleep(1)
                
                fresh_buttons = self.find_page_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element_safe(target):
                    time.sleep(2.5)
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{page_name}的{btn_text}功能")
                    self.close_dialog()
                    time.sleep(1)
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def get_all_menu_links(self):
        """获取所有菜单链接"""
        try:
            time.sleep(2)
            
            # 展开所有子菜单
            self.driver.execute_script("""
                var innerMenus = document.querySelectorAll('.lpx-inner-menu');
                innerMenus.forEach(function(menu) {
                    menu.classList.remove('collapsed');
                    menu.style.display = 'block';
                });
            """)
            time.sleep(1)
            
            # 查找所有菜单链接
            all_links = self.driver.find_elements(By.XPATH, 
                "//*[contains(@class, 'lpx-menu-item-link')]")
            
            menus = []
            for link in all_links:
                try:
                    text = link.text.strip()
                    href = link.get_attribute('href')
                    if text and len(text) < 100 and href:
                        menus.append({'text': text, 'href': href, 'element': link})
                except:
                    pass
            
            # 去重
            seen = set()
            unique = []
            for menu in menus:
                key = f"{menu['text']}_{menu['href']}"
                if key not in seen:
                    seen.add(key)
                    unique.append(menu)
            
            return unique
        except:
            return []
    
    def explore_by_function_modules(self):
        """按功能模块遍历"""
        print("\n" + "=" * 80)
        print("开始按功能模块遍历票务系统...")
        print("=" * 80)
        
        # 获取所有可用菜单
        all_menus = self.get_all_menu_links()
        print(f"\n✓ 系统中共找到 {len(all_menus)} 个菜单项")
        print("\n所有菜单项:")
        for i, menu in enumerate(all_menus, 1):
            print(f"  [{i}] {menu['text']}")
        print()
        
        # 按功能模块组织截图
        for module_name, module_info in self.function_modules.items():
            try:
                print(f"\n{'='*80}")
                print(f"功能模块: {module_name}")
                print(f"说明: {module_info['description']}")
                print(f"{'='*80}")
                
                # 查找匹配的菜单项
                matched_menus = []
                for menu in all_menus:
                    for keyword in module_info['menu_keywords']:
                        if keyword in menu['text']:
                            if menu not in matched_menus:
                                matched_menus.append(menu)
                                break
                
                if matched_menus:
                    print(f"  → 找到 {len(matched_menus)} 个相关菜单")
                    
                    for i, menu in enumerate(matched_menus, 1):
                        try:
                            print(f"\n  [{i}/{len(matched_menus)}] 菜单: {menu['text']}")
                            
                            # 返回主页
                            self.driver.get(self.base_url)
                            time.sleep(2)
                            
                            # 重新查找并点击菜单
                            fresh_menus = self.get_all_menu_links()
                            target = None
                            for fm in fresh_menus:
                                if fm['text'] == menu['text'] and fm['href'] == menu['href']:
                                    target = fm['element']
                                    break
                            
                            if target:
                                self.click_element_safe(target)
                                time.sleep(3)
                                
                                # 截图
                                self.save_screenshot(module_name, menu['text'], 
                                                   f"{module_name} - {menu['text']}")
                                
                                # 探索页面功能
                                self.explore_page(module_name, menu['text'])
                            
                        except Exception as e:
                            print(f"  ✗ 处理菜单出错: {str(e)}")
                else:
                    print(f"  → 未找到匹配的菜单")
                
            except Exception as e:
                print(f"✗ 处理功能模块出错: {str(e)}")
        
        print("\n" + "=" * 80)
        print(f"✓ 遍历完成！共截图 {len(self.screenshots)} 个功能点")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, function_modules, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading('票务系统 - 功能模块截图对照文档', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('票务系统（亿思维智慧门店）\n')
    info.add_run('系统地址：').bold = True
    info.add_run('https://xstest.axioxio.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y年%m月%d日 %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张\n')
    info.add_run('语言版本：').bold = True
    info.add_run('简体中文')
    
    # 添加功能模块说明
    doc.add_heading('功能模块说明', 1)
    for i, (module_name, module_info) in enumerate(function_modules.items(), 1):
        p = doc.add_paragraph()
        p.add_run(f'{i}. {module_name}：').bold = True
        p.add_run(module_info['description'])
    
    doc.add_paragraph()
    doc.add_heading('功能截图详情', 1)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    header = table.rows[0].cells
    header[0].text = '序号'
    header[1].text = '功能模块 / 功能点'
    header[2].text = '截图'
    
    for cell in header:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    widths = (Inches(0.5), Inches(2.5), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    for ss in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        cells[0].text = str(ss['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"【{ss['module']}】\n")
        run.font.bold = True
        run.font.size = Pt(10)
        
        run = desc_para.add_run(f"{ss['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{ss['description']}")
        run.font.size = Pt(8)
        
        img_path = os.path.join('/workspace/screenshots_by_modules', ss['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片加载失败]')
    
    doc.save(output_path)
    print(f"✓ Word文档: {output_path} ({os.path.getsize(output_path) / 1024 / 1024:.2f} MB)")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    
    print("=" * 80)
    print("票务系统功能模块截图工具")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = FunctionModuleScreenshotter(base_url, username, password)
        
        if screenshotter.login():
            screenshotter.explore_by_function_modules()
            
            json_path = os.path.join('screenshots_by_modules', 'index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'total': len(screenshotter.screenshots),
                    'language': '简体中文',
                    'function_modules': screenshotter.function_modules,
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            doc_path = '/workspace/票务系统功能模块截图对照文档.docx'
            create_word_document(screenshotter.screenshots, screenshotter.function_modules, doc_path)
            
            print("\n正在打包...")
            os.system("cd /workspace && rm -f 票务系统功能模块截图包.zip && " +
                     "zip -q -r 票务系统功能模块截图包.zip screenshots_by_modules/ 票务系统功能模块截图对照文档.docx")
            
            print("\n" + "=" * 80)
            print("✅ 完成！")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"🌐 语言版本: 简体中文")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统功能模块截图包.zip")
            print("=" * 80)
        else:
            print("\n❌ 登录失败")
    
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
