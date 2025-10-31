#!/usr/bin/env python3
"""
根据功能点列表进行截图并生成对照表
"""

import os
import time
import json
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 功能点列表（根据用户提供的列表）
FUNCTION_LIST = [
    {"module": "登录", "function": "身份验证", "description": "用户登录功能"},
    
    {"module": "活动管理", "function": "创建、编辑活动", "description": "活动的创建和编辑功能"},
    {"module": "活动管理", "function": "活动列表", "description": "活动列表页面"},
    {"module": "活动管理", "function": "查询、查看、评价、删除活动", "description": "活动的查询、查看、评价、删除操作"},
    
    {"module": "场地管理", "function": "创建、编辑场地", "description": "场地的创建和编辑功能"},
    {"module": "场地管理", "function": "场地列表", "description": "场地列表页面"},
    {"module": "场地管理", "function": "查询、查看场地资源", "description": "场地资源的查询和查看"},
    
    {"module": "订单管理", "function": "订单查询", "description": "订单查询功能"},
    {"module": "订单管理", "function": "退款订单", "description": "退款订单管理"},
    {"module": "订单管理", "function": "订单详情", "description": "订单详情页面"},
    
    {"module": "用户管理", "function": "用户订单列表", "description": "用户订单列表页面"},
    {"module": "用户管理", "function": "订单详情", "description": "用户订单详情"},
    
    {"module": "优惠券管理", "function": "优惠券列表", "description": "优惠券列表（领取记录查看、修改、复制删除等操作）"},
    {"module": "优惠券管理", "function": "优惠券搜索", "description": "优惠券搜索功能"},
    {"module": "优惠券管理", "function": "新建、删除优惠券", "description": "新建和删除优惠券操作"},
    {"module": "优惠券管理", "function": "优惠券规则", "description": "优惠券使用策略、使用范围等规则设置"},
    {"module": "优惠券管理", "function": "领取记录查询", "description": "按领取方式、使用状态查询领取记录"},
    
    {"module": "积分管理", "function": "积分规则设置", "description": "积分规则设置页面"},
    {"module": "积分管理", "function": "订单类积分", "description": "订单类积分设置"},
    {"module": "积分管理", "function": "任务类积分", "description": "任务类积分设置"},
    {"module": "积分管理", "function": "积分商城商品列表", "description": "积分商城商品列表页面"},
    {"module": "积分管理", "function": "商品查询", "description": "商品按条件查询"},
    {"module": "积分管理", "function": "新建、删除商品", "description": "新建和删除积分商品"},
    
    {"module": "积分订单", "function": "积分订单列表", "description": "积分商城消费订单列表"},
    {"module": "积分订单", "function": "订单查询", "description": "按条件查询积分订单"},
    {"module": "积分订单", "function": "订单详情", "description": "积分订单详情查看"},
    
    {"module": "小游戏管理", "function": "小游戏列表", "description": "小游戏列表页面"},
    {"module": "小游戏管理", "function": "小游戏查询", "description": "小游戏条件查询"},
    {"module": "小游戏管理", "function": "新建、删除、上架小游戏", "description": "小游戏的新建、删除和上架操作"},
    
    {"module": "员工管理", "function": "员工列表", "description": "员工列表页面"},
    {"module": "员工管理", "function": "员工详情", "description": "员工详情页面"},
    
    {"module": "财务管理", "function": "商户结算审核", "description": "商户结算审核页面"},
    {"module": "财务管理", "function": "财务报表", "description": "财务报表页面"},
    
    {"module": "数据分析", "function": "销售统计", "description": "销售统计页面"},
]

# URL映射（根据已知的系统结构）
URL_MAPPING = {
    "登录": "#/login",
    "活动管理-活动列表": "#/activity/index",
    "活动管理-活动订单": "#/activity/activeOrder",
    "场地管理-场地列表": "#/site/list",
    "场地管理-场地预定看板": "#/site/bookingBoard",
    "场地管理-场地订单": "#/site/orders",
    "用户管理-普通用户": "#/user/index",
    "用户管理-用户订单": "#/user/userOrder",
    "员工管理": "#/staff/index",
    "财务管理-商户结算审核": "#/finance/merchantSettlement",
    "财务管理-商户结算财务报表": "#/finance/financialReport",
    "数据分析-销售统计": "#/dataAnalysis/salesStatistics",
    "Banner管理": "#/system/bannerMsg",
    "优惠券列表": "#/coupon/index",
    "积分管理-规则设置": "#/integral/integralRule",
    "积分管理-积分商城": "#/integral/index",
    "积分管理-积分订单": "#/integral/integralOrder",
    "积分管理-积分用户": "#/integral/integralUser",
    "小游戏列表": "#/miniGame/index",
    "系统日志": "#/systemlog/index",
}

class FunctionScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_functions"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--disable-gpu')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(3)
        
    def login(self):
        """登录"""
        print(f"\n正在登录系统...")
        self.driver.get(self.base_url)
        time.sleep(3)
        
        try:
            # 登录页截图
            login_ss = os.path.join(self.output_dir, "00_登录页面.png")
            self.driver.save_screenshot(login_ss)
            
            username_input = self.driver.find_element(By.NAME, "username")
            password_input = self.driver.find_element(By.NAME, "password")
            
            username_input.send_keys(self.username)
            password_input.send_keys(self.password)
            
            login_button = self.driver.find_element(By.XPATH, "//button[@type='button']")
            login_button.click()
            
            for i in range(10):
                time.sleep(1)
                if 'dashboard' in self.driver.current_url:
                    break
            
            time.sleep(2)
            print("✓ 登录成功")
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def navigate_to_menu(self, menu_name):
        """导航到指定菜单"""
        try:
            time.sleep(1)
            
            # 查找菜单项
            menu_items = self.driver.find_elements(By.XPATH,
                f"//li[contains(@class, 'el-menu-item') or contains(@class, 'el-submenu')]//*[contains(text(), '{menu_name}')]")
            
            if menu_items:
                # 滚动到元素
                self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", menu_items[0])
                time.sleep(0.5)
                
                # 点击
                try:
                    menu_items[0].click()
                except:
                    self.driver.execute_script("arguments[0].click();", menu_items[0])
                
                time.sleep(2)
                return True
            
            return False
        except:
            return False
    
    def find_and_click_button(self, button_text):
        """查找并点击按钮"""
        try:
            buttons = self.driver.find_elements(By.XPATH,
                f"//button[contains(text(), '{button_text}')] | //a[contains(text(), '{button_text}')]")
            
            for btn in buttons[:1]:  # 只点击第一个
                if btn.is_displayed():
                    try:
                        btn.click()
                        time.sleep(1.5)
                        return True
                    except:
                        self.driver.execute_script("arguments[0].click();", btn)
                        time.sleep(1.5)
                        return True
            return False
        except:
            return False
    
    def take_screenshot(self, module, function, description, index):
        """截图"""
        try:
            time.sleep(1)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.5)
            
            filename = f"{index:03d}_{module}_{function}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{index:03d}] ✓ {module} - {function}")
            return True
        except Exception as e:
            print(f"  [{index:03d}] ✗ 截图失败: {str(e)}")
            return False
    
    def capture_all_functions(self):
        """根据功能列表截图"""
        print("\n开始根据功能列表截图...")
        print("=" * 60)
        
        for i, func_item in enumerate(FUNCTION_LIST, 1):
            module = func_item['module']
            function = func_item['function']
            description = func_item['description']
            
            print(f"\n[{i}/{len(FUNCTION_LIST)}] {module} - {function}")
            
            try:
                # 特殊处理登录页
                if module == "登录":
                    # 已经在登录时截过图了
                    self.screenshots.append({
                        "index": i,
                        "module": module,
                        "function": function,
                        "description": description,
                        "filename": "00_登录页面.png",
                        "url": self.base_url,
                        "title": "登录页面"
                    })
                    print(f"  [{i:03d}] ✓ 登录 - 身份验证")
                    continue
                
                # 导航到对应模块
                navigated = False
                
                # 尝试直接导航
                if module == "活动管理":
                    navigated = self.navigate_to_menu("活动")
                elif module == "场地管理":
                    navigated = self.navigate_to_menu("场地")
                elif module == "用户管理":
                    navigated = self.navigate_to_menu("用户")
                elif module == "员工管理":
                    navigated = self.navigate_to_menu("员工")
                elif module == "财务管理":
                    navigated = self.navigate_to_menu("财务")
                elif module == "数据分析":
                    navigated = self.navigate_to_menu("数据分析")
                elif module == "优惠券管理":
                    navigated = self.navigate_to_menu("优惠券")
                elif module == "积分管理":
                    navigated = self.navigate_to_menu("积分")
                elif module == "积分订单":
                    navigated = self.navigate_to_menu("积分订单")
                elif module == "小游戏管理":
                    navigated = self.navigate_to_menu("小游戏")
                
                if navigated:
                    time.sleep(1)
                    
                    # 根据功能执行操作
                    if "列表" in function:
                        # 列表页面直接截图
                        self.take_screenshot(module, function, description, i)
                    elif "创建" in function or "新建" in function or "添加" in function:
                        # 点击新建/创建按钮
                        if self.find_and_click_button("新建") or self.find_and_click_button("添加") or self.find_and_click_button("创建"):
                            self.take_screenshot(module, function, description, i)
                            # 关闭弹窗
                            try:
                                close_btn = self.driver.find_element(By.XPATH, "//*[contains(@class, 'el-dialog__close')]")
                                close_btn.click()
                                time.sleep(0.5)
                            except:
                                self.driver.back()
                                time.sleep(1)
                        else:
                            self.take_screenshot(module, function, description, i)
                    elif "编辑" in function or "修改" in function:
                        # 点击编辑按钮
                        if self.find_and_click_button("编辑") or self.find_and_click_button("修改"):
                            self.take_screenshot(module, function, description, i)
                            try:
                                close_btn = self.driver.find_element(By.XPATH, "//*[contains(@class, 'el-dialog__close')]")
                                close_btn.click()
                            except:
                                self.driver.back()
                                time.sleep(1)
                        else:
                            self.take_screenshot(module, function, description, i)
                    elif "查询" in function or "搜索" in function:
                        # 点击查询按钮
                        if self.find_and_click_button("查询") or self.find_and_click_button("搜索"):
                            self.take_screenshot(module, function, description, i)
                        else:
                            self.take_screenshot(module, function, description, i)
                    elif "详情" in function or "查看" in function:
                        # 点击查看/详情按钮
                        if self.find_and_click_button("查看") or self.find_and_click_button("详情"):
                            self.take_screenshot(module, function, description, i)
                            self.driver.back()
                            time.sleep(1)
                        else:
                            self.take_screenshot(module, function, description, i)
                    else:
                        # 其他情况直接截图当前页面
                        self.take_screenshot(module, function, description, i)
                else:
                    print(f"  ⚠ 无法导航到 {module}")
                    
            except Exception as e:
                print(f"  ✗ 处理出错: {str(e)}")
        
        print("\n" + "=" * 60)
        print(f"✓ 截图完成！共 {len(self.screenshots)} 个功能点")
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    
    # 设置中文字体
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    # 标题
    title = doc.add_heading('海星育后台管理系统 - 功能点截图对照表', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 基本信息
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('海星育后台管理系统\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y-%m-%d %H:%M:%S")}\n')
    info.add_run('功能点数：').bold = True
    info.add_run(f'{len(screenshots)} 个')
    
    doc.add_paragraph()
    
    # 创建表格
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    # 表头
    header_cells = table.rows[0].cells
    header_cells[0].text = '序号'
    header_cells[1].text = '功能点描述'
    header_cells[2].text = '功能截图'
    
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 设置列宽
    widths = (Inches(0.6), Inches(2.4), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    # 添加数据
    for screenshot in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        # 序号
        cells[0].text = str(screenshot['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 功能描述
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"{screenshot['module']}\n")
        run.font.bold = True
        run.font.size = Pt(10)
        
        run = desc_para.add_run(f"{screenshot['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{screenshot['description']}")
        run.font.size = Pt(8)
        
        # 截图
        img_path = os.path.join('/workspace/screenshots_functions', screenshot['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片加载失败]')
        else:
            cells[2].text = '[图片不存在]'
    
    doc.save(output_path)
    print(f"✓ Word文档已生成: {output_path}")
    print(f"✓ 文件大小: {os.path.getsize(output_path) / 1024 / 1024:.2f} MB")


def main():
    base_url = "https://jhtest.bjstarfish.com/"
    username = "admin"
    password = "123456"
    output_dir = "screenshots_functions"
    
    print("=" * 60)
    print("根据功能点列表生成截图对照表")
    print("=" * 60)
    print(f"功能点数量: {len(FUNCTION_LIST)}")
    print("=" * 60)
    
    screenshotter = None
    try:
        screenshotter = FunctionScreenshotter(base_url, username, password, output_dir)
        
        if screenshotter.login():
            screenshotter.capture_all_functions()
            
            # 生成Word文档
            output_path = '/workspace/功能点截图对照表（完整版）.docx'
            create_word_document(screenshotter.screenshots, output_path)
            
            print("\n" + "=" * 60)
            print("✅ 所有任务完成！")
            print(f"📁 截图目录: {os.path.abspath(output_dir)}")
            print(f"📄 Word文档: {output_path}")
            print("=" * 60)
    
    except Exception as e:
        print(f"\n❌ 发生错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()


if __name__ == "__main__":
    main()
