#!/usr/bin/env python3
"""
票务系统功能截图工具 - 按功能模块
"""

import os
import time
import json
import hashlib
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


class TicketSystemFunctionScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_ticket_functions"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_states = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        chrome_options.add_argument('--lang=zh-CN')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(3)
    
    def get_page_state_hash(self):
        """获取页面状态哈希"""
        try:
            url = self.driver.current_url
            title = self.driver.title
            try:
                main_text = self.driver.find_element(By.TAG_NAME, "main").text[:200]
            except:
                main_text = self.driver.find_element(By.TAG_NAME, "body").text[:200]
            state = f"{url}_{title}_{main_text}"
            return hashlib.md5(state.encode()).hexdigest()
        except:
            return None
    
    def switch_to_chinese(self):
        """切换到简体中文"""
        try:
            print("  → 切换到简体中文...")
            lang_button = self.driver.find_element(By.XPATH, "//button[contains(@class, 'dropdown-toggle')]")
            lang_button.click()
            time.sleep(1)
            chinese_link = self.driver.find_element(By.XPATH, "//a[contains(text(), '简体中文')]")
            chinese_link.click()
            time.sleep(3)
            print("  ✓ 已切换到简体中文")
            return True
        except:
            return False
    
    def login(self):
        """登录"""
        print(f"\n正在登录票务系统...")
        self.driver.get(self.base_url)
        time.sleep(5)
        
        try:
            self.switch_to_chinese()
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面")
            
            username_input = self.driver.find_element(By.NAME, "LoginInput.UserNameOrEmailAddress")
            password_input = self.driver.find_element(By.NAME, "LoginInput.Password")
            
            username_input.send_keys(self.username)
            password_input.send_keys(self.password)
            password_input.send_keys(Keys.RETURN)
            
            print("✓ 已提交登录")
            
            for i in range(15):
                time.sleep(1)
                if 'login' not in self.driver.current_url.lower():
                    print(f"✓ 登录成功！")
                    break
            
            time.sleep(3)
            self.save_screenshot("系统主页", "首页Dashboard", "登录后主页")
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            state_hash = self.get_page_state_hash()
            if state_hash and state_hash in self.visited_states:
                return False
            if state_hash:
                self.visited_states.add(state_hash)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.8)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def click_element_safe(self, element):
        """安全点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
            except:
                self.driver.execute_script("arguments[0].click();", element)
            return True
        except:
            return False
    
    def close_dialog(self):
        """关闭弹窗"""
        try:
            for selector in ["//button[contains(text(), '取消')]", "//button[contains(text(), '关闭')]",
                           "//button[contains(@class, 'close')]"]:
                try:
                    btns = self.driver.find_elements(By.XPATH, selector)
                    for btn in btns:
                        if btn.is_displayed():
                            btn.click()
                            time.sleep(1)
                            return
                except:
                    pass
            from selenium.webdriver.common.action_chains import ActionChains
            ActionChains(self.driver).send_keys(Keys.ESCAPE).perform()
            time.sleep(1)
        except:
            pass
    
    def find_menu_by_text(self, text):
        """根据文本查找菜单"""
        try:
            # 展开所有子菜单
            self.driver.execute_script("""
                var innerMenus = document.querySelectorAll('.lpx-inner-menu');
                innerMenus.forEach(function(menu) {
                    menu.classList.remove('collapsed');
                    menu.style.display = 'block';
                });
            """)
            time.sleep(1)
            
            # 查找包含文本的菜单项
            menu_link = self.driver.find_element(By.XPATH, 
                f"//a[contains(@class, 'lpx-menu-item-link') and contains(., '{text}')]")
            return menu_link
        except:
            return None
    
    def find_page_buttons(self):
        """查找页面按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH, 
                "//button | //a[contains(@class, 'btn')]")
            
            keywords = ['新增', '新建', '添加', '编辑', '修改', '查看', '详情', 
                       '删除', '设置', '配置', '导出', '导入', '审核', '发布', 
                       '提交', '保存', '搜索', '查询', '筛选']
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text and len(text) < 50 and any(k in text for k in keywords):
                            buttons.append({'element': btn, 'text': text})
                except:
                    pass
            
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique[:10]
        except:
            return []
    
    def explore_page_functions(self, module_name, page_name):
        """探索页面功能"""
        buttons = self.find_page_buttons()
        if buttons:
            print(f"      发现功能按钮: {len(buttons)} 个")
        
        for i, btn_info in enumerate(buttons, 1):
            try:
                btn_text = btn_info['text']
                print(f"        [{i}] 点击: {btn_text}")
                time.sleep(1)
                
                fresh_buttons = self.find_page_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element_safe(target):
                    time.sleep(2.5)
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{page_name}的{btn_text}功能")
                    self.close_dialog()
                    time.sleep(1)
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def capture_function_module(self, module_name, description, menu_keywords):
        """捕获功能模块的截图"""
        print(f"\n{'='*80}")
        print(f"功能模块: {module_name}")
        print(f"描述: {description}")
        print(f"{'='*80}")
        
        for keyword in menu_keywords:
            try:
                print(f"\n  → 查找菜单: {keyword}")
                
                # 返回主页
                self.driver.get(self.base_url)
                time.sleep(2)
                
                # 查找菜单
                menu_link = self.find_menu_by_text(keyword)
                
                if menu_link:
                    print(f"  ✓ 找到菜单: {keyword}")
                    
                    if self.click_element_safe(menu_link):
                        time.sleep(3)
                        
                        # 截图主页面
                        self.save_screenshot(module_name, keyword, f"{module_name} - {keyword}")
                        
                        # 探索页面功能
                        self.explore_page_functions(module_name, keyword)
                else:
                    print(f"  ✗ 未找到菜单: {keyword}")
                    
            except Exception as e:
                print(f"  ✗ 处理菜单 {keyword} 出错: {str(e)}")
    
    def capture_all_functions(self):
        """捕获所有功能模块"""
        print("\n" + "=" * 80)
        print("开始按功能模块截图...")
        print("=" * 80)
        
        # 定义功能模块和对应的菜单关键词
        function_modules = [
            {
                "name": "票务管理",
                "description": "实现对雪票电子或纸质的全面管理",
                "menus": ["票务管理", "电子券", "纸制票", "初始化", "电子/纸制票"]
            },
            {
                "name": "订单管理",
                "description": "对各类雪票、教务订单、租赁订单进行管理",
                "menus": ["订单", "雪票订单", "教务订单"]
            },
            {
                "name": "工作台",
                "description": "实现对现场售票及教练预约的下单管理",
                "menus": ["工作台", "收银", "发卡", "还卡", "团队下单"]
            },
            {
                "name": "通行管理",
                "description": "通过闸机的管控实现对进出雪场的人员进行管理",
                "menus": ["闸机", "人脸库", "通行记录", "通行规则", "卡片管理", "软件版本"]
            },
            {
                "name": "控制面板",
                "description": "包括对雪票产品设置、电子/纸质票设置、闸机配置等运营参数的管理",
                "menus": ["产品管理", "产品分类", "小程序", "分销商", "租赁"]
            },
            {
                "name": "报表管理",
                "description": "提供各类运营报表",
                "menus": ["报表", "核销报表", "数据看板"]
            },
            {
                "name": "系统管理",
                "description": "系统基础功能的设置",
                "menus": ["管理", "系统设置", "日志", "文件"]
            }
        ]
        
        for module in function_modules:
            self.capture_function_module(
                module["name"],
                module["description"],
                module["menus"]
            )
        
        print("\n" + "=" * 80)
        print(f"✓ 截图完成！共截图 {len(self.screenshots)} 个功能点")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading('票务系统 - 功能截图对照文档', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('票务系统（亿思维智慧门店）\n')
    info.add_run('系统地址：').bold = True
    info.add_run('https://xstest.axioxio.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y年%m月%d日 %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张')
    
    doc.add_heading('功能模块说明', 1)
    
    modules_desc = [
        ('票务管理', '实现对雪票电子或纸质的全面管理'),
        ('订单管理', '对各类雪票、教务订单、租赁订单进行管理'),
        ('工作台', '实现对现场售票及教练预约的下单管理'),
        ('通行管理', '通过闸机的管控实现对进出雪场的人员进行管理；通行方式可以是人脸或雪票二维码'),
        ('控制面板', '包括对雪票产品设置，电子/纸质票设置、闸机配置、售票小程序、教练、分销商、节假日、租赁物等运营参数的管理'),
        ('报表管理', '提供各类运营报表，包括票务核销报表、教学核销报表、教学数据看板等'),
        ('系统管理', '系统基础功能的设置，包括文件，身份标识，日志及相关系统参数的设置')
    ]
    
    for mod_name, mod_desc in modules_desc:
        p = doc.add_paragraph()
        p.add_run(f'{mod_name}：').bold = True
        p.add_run(mod_desc)
    
    doc.add_paragraph()
    doc.add_heading('功能截图详情', 1)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    header = table.rows[0].cells
    header[0].text = '序号'
    header[1].text = '功能模块 / 功能点'
    header[2].text = '截图'
    
    for cell in header:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    widths = (Inches(0.5), Inches(2.5), Inches(4.0))
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
    
    for ss in screenshots:
        row = table.add_row()
        cells = row.cells
        
        for idx, width in enumerate(widths):
            cells[idx].width = width
        
        cells[0].text = str(ss['index'])
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        desc_para = cells[1].paragraphs[0]
        run = desc_para.add_run(f"【{ss['module']}】\n")
        run.font.bold = True
        run.font.size = Pt(10)
        
        run = desc_para.add_run(f"{ss['function']}\n\n")
        run.font.size = Pt(9)
        
        run = desc_para.add_run(f"{ss['description']}")
        run.font.size = Pt(8)
        
        img_path = os.path.join('/workspace/screenshots_ticket_functions', ss['filename'])
        if os.path.exists(img_path):
            img_para = cells[2].paragraphs[0]
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = img_para.add_run()
                run.add_picture(img_path, width=Inches(3.8))
            except:
                img_para.add_run('[图片加载失败]')
    
    doc.save(output_path)
    print(f"✓ Word文档: {output_path} ({os.path.getsize(output_path) / 1024 / 1024:.2f} MB)")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    
    print("=" * 80)
    print("票务系统功能截图工具（按功能模块）")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = TicketSystemFunctionScreenshotter(base_url, username, password)
        
        if screenshotter.login():
            screenshotter.capture_all_functions()
            
            json_path = os.path.join('screenshots_ticket_functions', 'index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'total': len(screenshotter.screenshots),
                    'language': '简体中文',
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            doc_path = '/workspace/票务系统功能截图对照文档.docx'
            create_word_document(screenshotter.screenshots, doc_path)
            
            print("\n正在打包...")
            os.system("cd /workspace && rm -f 票务系统功能截图包.zip && " +
                     "zip -q -r 票务系统功能截图包.zip screenshots_ticket_functions/ 票务系统功能截图对照文档.docx")
            
            zip_size = os.path.getsize('/workspace/票务系统功能截图包.zip') / 1024 / 1024
            
            print("\n" + "=" * 80)
            print("✅ 完成！")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"🌐 语言版本: 简体中文")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统功能截图包.zip ({zip_size:.2f} MB)")
            print("=" * 80)
            
            # 提供下载链接说明
            print("\n下载方式：")
            print("  请下载文件：/workspace/票务系统功能截图包.zip")
        else:
            print("\n❌ 登录失败")
    
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
