#!/usr/bin/env python3
"""
票务系统功能截图工具 - 按功能结构
"""

import os
import time
import json
import hashlib
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


# 功能结构定义
FUNCTION_STRUCTURE = {
    "票务管理": {
        "description": "实现对雪票电子或纸质的全面管理",
        "keywords": ["票务", "电子券", "纸制票", "电子票", "票务管理", "雪票"]
    },
    "订单管理": {
        "description": "对各类雪票、教务订单、租赁订单进行管理",
        "keywords": ["订单", "订单管理", "票务管理"]
    },
    "工作台": {
        "description": "实现对现场售票及教练预约的下单管理",
        "keywords": ["工作台", "收银", "发卡", "还卡", "团队下单", "预约"]
    },
    "通行管理": {
        "description": "通过闸机的管控实现对进出雪场的人员进行管理；通行方式可以是人脸或雪票二维码",
        "keywords": ["闸机", "人脸", "通行", "卡片", "通行记录", "通行规则"]
    },
    "控制面板": {
        "description": "包括对雪票产品设置，电子/纸质票设置、闸机配置、售票小程序、教练、分销商、节假日、租赁物等运营参数的管理",
        "keywords": ["产品", "分类", "小程序", "分销商", "租赁", "设置", "配置"]
    },
    "报表管理": {
        "description": "提供各类运营报表，包括票务核销报表、教学核销报表、教学数据看板等",
        "keywords": ["报表", "统计", "分析", "数据"]
    },
    "系统管理": {
        "description": "系统基础功能的设置，包括文件，身份标识，日志及相关系统参数的设置",
        "keywords": ["管理", "系统", "用户", "角色", "权限", "日志", "设置"]
    }
}


class FunctionStructureScreenshotter:
    def __init__(self, base_url, username, password, output_dir="screenshots_by_structure"):
        self.base_url = base_url
        self.username = username
        self.password = password
        self.output_dir = output_dir
        self.screenshots = []
        self.visited_states = set()
        self.screenshot_index = 1
        
        os.makedirs(output_dir, exist_ok=True)
        
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--no-sandbox')
        chrome_options.add_argument('--disable-dev-shm-usage')
        chrome_options.add_argument('--window-size=1920,1080')
        chrome_options.add_argument('--ignore-certificate-errors')
        chrome_options.add_argument('--lang=zh-CN')
        
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.implicitly_wait(2)
    
    def get_page_state_hash(self):
        """获取页面状态哈希"""
        try:
            url = self.driver.current_url
            title = self.driver.title
            try:
                main_text = self.driver.find_element(By.TAG_NAME, "main").text[:200]
            except:
                main_text = self.driver.find_element(By.TAG_NAME, "body").text[:200]
            state = f"{url}_{title}_{main_text}"
            return hashlib.md5(state.encode()).hexdigest()
        except:
            return None
    
    def switch_to_chinese(self):
        """切换到简体中文"""
        try:
            print("  → 切换到简体中文...")
            lang_button = self.driver.find_element(By.XPATH, "//button[contains(@class, 'dropdown-toggle')]")
            lang_button.click()
            time.sleep(1)
            chinese_link = self.driver.find_element(By.XPATH, "//a[contains(text(), '简体中文')]")
            chinese_link.click()
            time.sleep(3)
            print("  ✓ 已切换到简体中文")
            return True
        except:
            return False
    
    def login(self):
        """登录"""
        print(f"\n正在登录票务系统...")
        self.driver.get(self.base_url)
        time.sleep(5)
        
        try:
            self.switch_to_chinese()
            self.save_screenshot("系统登录", "登录页面", "票务系统登录界面（中文）")
            
            username_input = self.driver.find_element(By.NAME, "LoginInput.UserNameOrEmailAddress")
            password_input = self.driver.find_element(By.NAME, "LoginInput.Password")
            
            username_input.send_keys(self.username)
            password_input.send_keys(self.password)
            password_input.send_keys(Keys.RETURN)
            
            print("✓ 已提交登录")
            
            for i in range(15):
                time.sleep(1)
                if 'login' not in self.driver.current_url.lower():
                    print(f"✓ 登录成功！")
                    break
            
            time.sleep(3)
            self.save_screenshot("系统主页", "首页Dashboard", "登录后主页")
            return True
        except Exception as e:
            print(f"✗ 登录失败: {str(e)}")
            return False
    
    def save_screenshot(self, module, function, description):
        """保存截图"""
        try:
            state_hash = self.get_page_state_hash()
            if state_hash and state_hash in self.visited_states:
                return False
            if state_hash:
                self.visited_states.add(state_hash)
            
            time.sleep(1.5)
            self.driver.execute_script("window.scrollTo(0, 0);")
            time.sleep(0.8)
            
            safe_name = "".join(c if c.isalnum() or c in (' ', '-', '_') else '_' for c in f"{module}_{function}")
            safe_name = safe_name[:70]
            
            filename = f"{self.screenshot_index:04d}_{safe_name}.png"
            filepath = os.path.join(self.output_dir, filename)
            
            self.driver.save_screenshot(filepath)
            
            self.screenshots.append({
                "index": self.screenshot_index,
                "module": module,
                "function": function,
                "description": description,
                "filename": filename,
                "url": self.driver.current_url,
                "title": self.driver.title
            })
            
            print(f"  [{self.screenshot_index:04d}] ✓ {module} - {function}")
            self.screenshot_index += 1
            return True
        except Exception as e:
            print(f"  ✗ 截图失败: {str(e)}")
            return False
    
    def click_element_safe(self, element):
        """安全点击元素"""
        try:
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
            time.sleep(0.5)
            try:
                element.click()
            except:
                self.driver.execute_script("arguments[0].click();", element)
            return True
        except:
            return False
    
    def close_dialog(self):
        """关闭弹窗"""
        try:
            for selector in ["//button[contains(text(), '取消')]", "//button[contains(text(), '关闭')]",
                           "//button[contains(@class, 'close')]", "//a[contains(text(), '返回')]"]:
                try:
                    btns = self.driver.find_elements(By.XPATH, selector)
                    for btn in btns:
                        if btn.is_displayed():
                            btn.click()
                            time.sleep(1)
                            return
                except:
                    pass
            from selenium.webdriver.common.action_chains import ActionChains
            ActionChains(self.driver).send_keys(Keys.ESCAPE).perform()
            time.sleep(1)
        except:
            pass
    
    def find_page_buttons(self):
        """查找页面按钮"""
        buttons = []
        try:
            button_elements = self.driver.find_elements(By.XPATH, 
                "//button | //a[contains(@class, 'btn')]")
            
            keywords = ['新增', '新建', '添加', '编辑', '修改', '查看', '详情', 
                       '删除', '设置', '配置', '导出', '导入', '审核', '发布', 
                       '提交', '保存', '搜索', '查询', '筛选']
            
            for btn in button_elements:
                try:
                    if btn.is_displayed() and btn.is_enabled():
                        text = btn.text.strip()
                        if text and len(text) < 50 and any(k in text for k in keywords):
                            buttons.append({'element': btn, 'text': text})
                except:
                    pass
            
            seen = set()
            unique = []
            for btn in buttons:
                if btn['text'] not in seen:
                    seen.add(btn['text'])
                    unique.append(btn)
            
            return unique[:10]
        except:
            return []
    
    def explore_page(self, module_name, page_name):
        """探索页面功能"""
        buttons = self.find_page_buttons()
        if buttons:
            print(f"      发现功能按钮: {len(buttons)} 个")
        
        for i, btn_info in enumerate(buttons, 1):
            try:
                btn_text = btn_info['text']
                print(f"        [{i}] 点击: {btn_text}")
                time.sleep(1)
                
                fresh_buttons = self.find_page_buttons()
                target = None
                for fb in fresh_buttons:
                    if fb['text'] == btn_text:
                        target = fb['element']
                        break
                
                if target and self.click_element_safe(target):
                    time.sleep(2.5)
                    function_name = f"{page_name} - {btn_text}"
                    self.save_screenshot(module_name, function_name, f"{page_name}的{btn_text}功能")
                    self.close_dialog()
                    time.sleep(1)
            except Exception as e:
                print(f"        ✗ 处理按钮出错: {str(e)}")
    
    def get_all_menu_items(self):
        """获取所有菜单项"""
        try:
            time.sleep(2)
            
            # 展开所有子菜单
            self.driver.execute_script("""
                var innerMenus = document.querySelectorAll('.lpx-inner-menu');
                innerMenus.forEach(function(menu) {
                    menu.classList.remove('collapsed');
                    menu.style.display = 'block';
                });
            """)
            time.sleep(1)
            
            # 查找所有菜单项
            menu_links = self.driver.find_elements(By.XPATH, "//a[contains(@class, 'lpx-menu-item-link')]")
            
            menus = []
            for link in menu_links:
                try:
                    text_elem = link.find_element(By.XPATH, ".//span[contains(@class, 'lpx-menu-item-text')]")
                    text = text_elem.text.strip()
                    href = link.get_attribute('href')
                    
                    if text and len(text) < 100:
                        # 判断是否在子菜单中
                        try:
                            parent_ul = link.find_element(By.XPATH, "ancestor::ul[contains(@class, 'lpx-inner-menu')]")
                            is_submenu = True
                        except:
                            is_submenu = False
                        
                        menus.append({
                            'text': text,
                            'href': href,
                            'is_submenu': is_submenu,
                            'element': link
                        })
                except:
                    pass
            
            return menus
        except Exception as e:
            print(f"获取菜单失败: {str(e)}")
            return []
    
    def categorize_menu_by_function(self, menu_text):
        """根据关键词将菜单归类到功能模块"""
        for function_name, function_info in FUNCTION_STRUCTURE.items():
            keywords = function_info['keywords']
            for keyword in keywords:
                if keyword in menu_text:
                    return function_name
        return "其他功能"
    
    def explore_all_functions(self):
        """按功能结构遍历"""
        print("\n" + "=" * 80)
        print("开始按功能结构遍历票务系统...")
        print("=" * 80)
        
        # 获取所有菜单
        all_menus = self.get_all_menu_items()
        print(f"\n✓ 发现 {len(all_menus)} 个菜单项")
        
        # 按功能结构分类
        categorized = {}
        for menu in all_menus:
            category = self.categorize_menu_by_function(menu['text'])
            if category not in categorized:
                categorized[category] = []
            categorized[category].append(menu)
        
        # 打印分类结果
        print("\n功能模块分类:")
        for function_name, menus in categorized.items():
            if function_name in FUNCTION_STRUCTURE:
                desc = FUNCTION_STRUCTURE[function_name]['description']
                print(f"\n【{function_name}】- {desc}")
            else:
                print(f"\n【{function_name}】")
            for menu in menus:
                submenu_mark = "  └─" if menu['is_submenu'] else "  ├─"
                print(f"{submenu_mark} {menu['text']}")
        
        # 按功能模块顺序遍历
        for function_name, function_info in FUNCTION_STRUCTURE.items():
            if function_name not in categorized:
                continue
            
            print(f"\n{'='*80}")
            print(f"功能模块: {function_name}")
            print(f"描述: {function_info['description']}")
            print(f"{'='*80}")
            
            menus = categorized[function_name]
            
            for i, menu in enumerate(menus, 1):
                try:
                    print(f"\n  [{i}/{len(menus)}] {menu['text']}")
                    
                    # 返回主页
                    self.driver.get(self.base_url)
                    time.sleep(2)
                    
                    # 展开所有子菜单
                    self.driver.execute_script("""
                        var innerMenus = document.querySelectorAll('.lpx-inner-menu');
                        innerMenus.forEach(function(menu) {
                            menu.classList.remove('collapsed');
                            menu.style.display = 'block';
                        });
                    """)
                    time.sleep(1)
                    
                    # 重新查找并点击菜单
                    try:
                        menu_link = self.driver.find_element(By.XPATH, 
                            f"//a[contains(@class, 'lpx-menu-item-link')]//span[contains(@class, 'lpx-menu-item-text') and contains(text(), '{menu['text']}')]/ancestor::a")
                        
                        if self.click_element_safe(menu_link):
                            time.sleep(3)
                            
                            # 截图
                            self.save_screenshot(function_name, menu['text'], 
                                               f"{function_name} - {menu['text']}")
                            
                            # 探索页面功能
                            self.explore_page(function_name, menu['text'])
                    except Exception as e:
                        print(f"    ✗ 点击菜单失败: {str(e)}")
                
                except Exception as e:
                    print(f"    ✗ 处理菜单出错: {str(e)}")
        
        print("\n" + "=" * 80)
        print(f"✓ 遍历完成！共截图 {len(self.screenshots)} 个功能点")
        print("=" * 80)
    
    def close(self):
        if self.driver:
            self.driver.quit()


def create_word_document(screenshots, output_path):
    """生成Word文档"""
    print("\n正在生成Word文档...")
    
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    title = doc.add_heading('票务系统 - 功能截图对照文档', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 系统信息
    info = doc.add_paragraph()
    info.add_run('系统名称：').bold = True
    info.add_run('票务系统（亿思维智慧门店）\n')
    info.add_run('系统地址：').bold = True
    info.add_run('https://xstest.axioxio.com/\n')
    info.add_run('生成时间：').bold = True
    info.add_run(f'{time.strftime("%Y年%m月%d日 %H:%M:%S")}\n')
    info.add_run('截图总数：').bold = True
    info.add_run(f'{len(screenshots)} 张')
    
    doc.add_paragraph()
    
    # 功能模块说明
    doc.add_heading('功能模块说明', level=1)
    for i, (function_name, function_info) in enumerate(FUNCTION_STRUCTURE.items(), 1):
        para = doc.add_paragraph()
        run = para.add_run(f"{i}. {function_name}：")
        run.bold = True
        para.add_run(function_info['description'])
    
    doc.add_page_break()
    
    # 按功能模块分组截图
    for function_name in FUNCTION_STRUCTURE.keys():
        function_screenshots = [s for s in screenshots if s['module'] == function_name]
        
        if not function_screenshots:
            continue
        
        # 功能模块标题
        doc.add_heading(f'{function_name}（{len(function_screenshots)}张）', level=1)
        doc.add_paragraph(FUNCTION_STRUCTURE[function_name]['description'])
        
        # 创建表格
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        
        header = table.rows[0].cells
        header[0].text = '序号'
        header[1].text = '功能点'
        header[2].text = '截图'
        
        for cell in header:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        widths = (Inches(0.5), Inches(2.0), Inches(4.5))
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
        
        # 添加截图
        for ss in function_screenshots:
            row = table.add_row()
            cells = row.cells
            
            for idx, width in enumerate(widths):
                cells[idx].width = width
            
            cells[0].text = str(ss['index'])
            cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            desc_para = cells[1].paragraphs[0]
            run = desc_para.add_run(f"{ss['function']}\n\n")
            run.font.size = Pt(9)
            run.font.bold = True
            
            run = desc_para.add_run(f"{ss['description']}")
            run.font.size = Pt(8)
            
            img_path = os.path.join('/workspace/screenshots_by_structure', ss['filename'])
            if os.path.exists(img_path):
                img_para = cells[2].paragraphs[0]
                img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    run = img_para.add_run()
                    run.add_picture(img_path, width=Inches(4.3))
                except:
                    img_para.add_run('[图片加载失败]')
        
        doc.add_page_break()
    
    doc.save(output_path)
    print(f"✓ Word文档: {output_path} ({os.path.getsize(output_path) / 1024 / 1024:.2f} MB)")


def main():
    base_url = "https://xstest.axioxio.com/"
    username = "13811458301"
    password = "4<z%0/RS"
    
    print("=" * 80)
    print("票务系统功能截图工具（按功能结构）")
    print("=" * 80)
    
    screenshotter = None
    try:
        screenshotter = FunctionStructureScreenshotter(base_url, username, password)
        
        if screenshotter.login():
            screenshotter.explore_all_functions()
            
            # 保存JSON
            json_path = os.path.join('screenshots_by_structure', 'index.json')
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump({
                    'total': len(screenshotter.screenshots),
                    'language': '简体中文',
                    'function_structure': FUNCTION_STRUCTURE,
                    'screenshots': screenshotter.screenshots
                }, f, ensure_ascii=False, indent=2)
            
            # 生成Word
            doc_path = '/workspace/票务系统功能截图对照文档.docx'
            create_word_document(screenshotter.screenshots, doc_path)
            
            # 打包
            print("\n正在打包...")
            os.system("cd /workspace && rm -f 票务系统功能截图包.zip && " +
                     "zip -q -r 票务系统功能截图包.zip screenshots_by_structure/ 票务系统功能截图对照文档.docx")
            
            print("\n" + "=" * 80)
            print("✅ 完成！")
            print(f"📊 截图总数: {len(screenshotter.screenshots)} 张")
            print(f"🌐 语言版本: 简体中文")
            print(f"📄 Word文档: {doc_path}")
            print(f"📦 压缩包: /workspace/票务系统功能截图包.zip")
            print("=" * 80)
        else:
            print("\n❌ 登录失败")
    
    except Exception as e:
        print(f"\n❌ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
    
    finally:
        if screenshotter:
            screenshotter.close()
            print("\n浏览器已关闭")


if __name__ == "__main__":
    main()
